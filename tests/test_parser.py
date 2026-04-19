from pathlib import Path

from docx import Document

from src.parser import (
    _extract_footnotes_from_xml,
    detect_summary_statistics,
    extract_raw_docx,
    normalize_extracted_content,
    parse_csv_schema,
    parse_docx_shell,
    parse_inputs,
)


def test_parse_docx_shell_extracts_table_blueprint(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    file_path = data_dir / "sample.docx"

    doc = Document()
    doc.add_paragraph("TLG Shell Title")
    doc.add_paragraph("Section One:")
    doc.add_paragraph("Population: Safety Population")
    doc.add_paragraph("Output format: JSON")
    table = doc.add_table(rows=9, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Age (yr)"
    table.cell(1, 1).text = ""
    table.cell(2, 0).text = "n"
    table.cell(2, 1).text = "100"
    table.cell(3, 0).text = "Mean (SD)"
    table.cell(3, 1).text = "12.3 (1.1)"
    table.cell(4, 0).text = "Median"
    table.cell(4, 1).text = "12.1"
    table.cell(5, 0).text = "Min-max"
    table.cell(5, 1).text = "10.0-15.0"
    table.cell(6, 0).text = "Sex"
    table.cell(6, 1).text = ""
    table.cell(7, 0).text = "Male"
    table.cell(7, 1).text = "60 (60%)"
    table.cell(8, 0).text = "Female"
    table.cell(8, 1).text = "40 (40%)"
    doc.save(file_path)

    result = parse_docx_shell("sample.docx", data_dir=str(data_dir))

    assert result["source_file"] == str(file_path)
    assert result["document"]["title"] == "TLG Shell Title"
    assert result["document"]["paragraphs"] == [
        "TLG Shell Title",
        "Section One:",
        "Population: Safety Population",
        "Output format: JSON",
    ]
    assert result["table_shells"][0]["columns"] == ["Value"]
    assert result["table_shells"][0]["group_sizes"] == {}
    assert result["table_shells"][0]["ordered_rows"] == [
        {"kind": "group_header", "label": "Age (yr)"},
        {"kind": "data_row", "group": "Age (yr)", "label": "n", "format_hint": None},
        {"kind": "data_row", "group": "Age (yr)", "label": "Mean (SD)", "format_hint": None},
        {"kind": "data_row", "group": "Age (yr)", "label": "Median", "format_hint": None},
        {"kind": "data_row", "group": "Age (yr)", "label": "Min–max", "format_hint": None},
        {"kind": "group_header", "label": "Sex"},
        {"kind": "data_row", "group": "Sex", "label": "Male", "format_hint": None},
        {"kind": "data_row", "group": "Sex", "label": "Female", "format_hint": None},
    ]
    assert result["table_shells"][0]["row_groups"] == [
        {"name": "Age (yr)", "subrows": ["n", "Mean (SD)", "Median", "Min–max"]},
        {"name": "Sex", "subrows": ["Male", "Female"]},
    ]


def test_parse_csv_schema_reads_header_only(tmp_path: Path) -> None:
    csv_path = tmp_path / "adsl.csv"
    csv_path.write_text("TRT01A,AGE,SEX\nA,42,M\n", encoding="utf-8")

    result = parse_csv_schema(str(csv_path))

    assert result == {
        "csv_path": str(csv_path),
        "dataset_name": "adsl",
        "columns": [
            {"name": "TRT01A", "normalized_name": "trt01a"},
            {"name": "AGE", "normalized_name": "age"},
            {"name": "SEX", "normalized_name": "sex"},
        ],
    }


def test_parse_inputs_combines_docx_and_csv_specs(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    docx_path = data_dir / "combined.docx"
    csv_path = data_dir / "combined.csv"

    doc = Document()
    doc.add_paragraph("Combined Title")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = ""
    table.cell(0, 1).text = "Arm A (N=10)"
    table.cell(1, 0).text = "Sex"
    table.cell(1, 1).text = ""
    table.cell(2, 0).text = "Male"
    table.cell(2, 1).text = "5"
    doc.save(docx_path)
    csv_path.write_text("TRT01A,SEX\nA,M\n", encoding="utf-8")

    result = parse_inputs("combined.docx", str(csv_path), data_dir=str(data_dir))

    assert result["docx_spec"]["document"]["title"] == "Combined Title"
    assert result["docx_spec"]["table_shells"][0]["columns"] == ["Arm A"]
    assert result["csv_schema"]["columns"][0]["name"] == "TRT01A"


def test_extract_and_normalize_raw_docx_layers(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    file_path = data_dir / "layered.docx"

    doc = Document()
    doc.add_paragraph("  Title  ")
    doc.add_paragraph("  ")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Col1"
    table.cell(0, 1).text = "Col2"
    table.cell(1, 0).text = "Section A"
    table.cell(1, 1).text = ""
    table.cell(2, 0).text = "- detail"
    table.cell(2, 1).text = "note"
    doc.save(file_path)

    raw = extract_raw_docx("layered.docx", data_dir=str(data_dir))
    assert raw["source_file"] == str(file_path)
    assert raw["paragraphs"][0] == "  Title  "

    normalized = normalize_extracted_content(raw)
    assert normalized["paragraphs"] == ["Title"]
    assert normalized["tables"][0][2][0] == "detail"


def test_detect_summary_statistics_from_row_labels() -> None:
    rows = [
        {"row_label": "n"},
        {"row_label": "Mean (SD)"},
        {"row_label": "Median"},
        {"row_label": "Min–max"},
    ]
    assert detect_summary_statistics(rows) == [
        "count",
        "mean",
        "sd",
        "median",
        "min",
        "max",
    ]


def test_row_label_normalization_for_newlines_and_symbols() -> None:
    normalized = {
        "source_file": "TLG_shell_t1.docx",
        "paragraphs": ["Demo Table", "Population: Safety Population"],
        "tables": [
            [
                ["", "Group 1\n(N=10)"],
                ["Age group (yr)", ""],
                ["≥65", "3"],
                ["65", "2"],
                ["Minmax", "1-9"],
                ["Race", ""],
                ["Native Hawaiian or \nother Pacific Islander", "2"],
            ]
        ],
        "footnotes": [],
    }

    parsed = normalize_extracted_content(normalized)
    labels = [row[0] for row in parsed["tables"][0][1:]]

    assert "≥65" in labels
    assert labels.count("≥65") >= 2
    assert "Min–max" in labels
    assert "Native Hawaiian or other Pacific Islander" in labels


def test_extract_footnotes_from_xml_blob() -> None:
    xml_blob = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:footnote w:id="1">
    <w:p><w:r><w:t>Footnote 1</w:t></w:r></w:p>
  </w:footnote>
</w:footnotes>
"""

    assert _extract_footnotes_from_xml(xml_blob) == ["Footnote 1"]
