"""Schema-only parser for DOCX shell and CSV metadata."""

from __future__ import annotations

import csv
import re
import zipfile
from io import BytesIO
from pathlib import Path
import xml.etree.ElementTree as ET
from typing import Any, Dict, List

from docx import Document


def _normalize_whitespace(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _normalize_row_label(label: str) -> str:
    normalized = _normalize_whitespace(label)
    normalized = re.sub(r"^\-\s*", "", normalized)
    if normalized == "65":
        return "≥65"
    if re.fullmatch(r"min\s*[-–—]?\s*max", normalized, flags=re.IGNORECASE):
        return "Min–max"
    return normalized


def _extract_table_blueprint(table: Any) -> Dict[str, Any]:
    raw_rows: List[List[str]] = []
    for row in table.rows:
        raw_rows.append([_normalize_whitespace(cell.text) for cell in row.cells])

    if not raw_rows:
        return {
            "columns": [],
            "group_sizes": {},
            "ordered_rows": [],
            "row_groups": [],
        }

    header = raw_rows[0]
    columns: List[str] = []
    group_sizes: Dict[str, str] = {}
    for raw in header[1:]:
        if not raw:
            continue
        match = re.match(r"^(.*?)\s*\(\s*N\s*=\s*([^)]+)\s*\)\s*$", raw, flags=re.IGNORECASE)
        if match:
            name = _normalize_whitespace(match.group(1))
            n_value = _normalize_whitespace(match.group(2))
        else:
            name = _normalize_whitespace(raw.splitlines()[0])
            n_value = ""
        if not name:
            continue
        columns.append(name)
        if n_value:
            group_sizes[name] = n_value

    ordered_rows: List[Dict[str, Any]] = []
    row_groups: List[Dict[str, Any]] = []
    current_group: str | None = None

    for row in raw_rows[1:]:
        label = _normalize_row_label(row[0] if row else "")
        if not label:
            continue

        values = [row[idx].strip() if idx < len(row) else "" for idx in range(1, len(columns) + 1)]

        if all(not value for value in values):
            current_group = label
            row_groups.append({"name": label, "subrows": []})
            ordered_rows.append({"kind": "group_header", "label": label})
            continue

        if current_group is None:
            current_group = "Ungrouped"
            row_groups.append({"name": current_group, "subrows": []})

        row_groups[-1]["subrows"].append(label)
        ordered_rows.append(
            {
                "kind": "data_row",
                "group": current_group,
                "label": label,
                "format_hint": None,
            }
        )

    return {
        "columns": columns,
        "group_sizes": group_sizes,
        "ordered_rows": ordered_rows,
        "row_groups": row_groups,
    }


def parse_docx_shell(filename: str, data_dir: str = "data") -> Dict[str, Any]:
    """Parse DOCX shell structure only (no dataset access)."""
    file_path = Path(data_dir) / filename
    document = Document(file_path)

    paragraphs = [_normalize_whitespace(p.text) for p in document.paragraphs]
    paragraphs = [p for p in paragraphs if p]
    title = paragraphs[0] if paragraphs else file_path.stem

    table_shells = [_extract_table_blueprint(table) for table in document.tables]

    return {
        "source_file": str(file_path),
        "document": {
            "title": title,
            "paragraphs": paragraphs,
        },
        "table_shells": table_shells,
    }


def parse_csv_schema(csv_path: str) -> Dict[str, Any]:
    """Read only CSV header (schema-only mode)."""
    path = Path(csv_path)
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader, [])

    columns: List[Dict[str, str]] = []
    for raw_name in header:
        name = _normalize_whitespace(raw_name)
        if not name:
            continue
        columns.append(
            {
                "name": name,
                "normalized_name": re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_"),
            }
        )

    return {
        "csv_path": str(path),
        "dataset_name": re.sub(r"[^a-z0-9]+", "_", path.stem.lower()).strip("_") or "dataset",
        "columns": columns,
    }


def parse_inputs(docx_filename: str, csv_path: str, data_dir: str = "data") -> Dict[str, Any]:
    """Combined schema parse for DOCX shell and CSV attributes."""
    docx_spec = parse_docx_shell(filename=docx_filename, data_dir=data_dir)
    csv_schema = parse_csv_schema(csv_path=csv_path)
    return {
        "docx_spec": docx_spec,
        "csv_schema": csv_schema,
    }


def _extract_footnotes_from_xml(xml_blob: bytes) -> List[str]:
    """Extract visible Word footnotes from a footnotes.xml blob."""
    if not xml_blob:
        return []

    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = ET.parse(BytesIO(xml_blob)).getroot()
    notes: List[str] = []
    for footnote in root.findall("w:footnote", namespace):
        footnote_type = footnote.attrib.get(f"{{{namespace['w']}}}type", "")
        if footnote_type:
            continue
        text = "".join(node.text or "" for node in footnote.findall(".//w:t", namespace))
        normalized = _normalize_whitespace(text)
        if normalized:
            notes.append(normalized)
    return notes


def extract_raw_docx(filename: str, data_dir: str = "data") -> Dict[str, Any]:
    """Extract raw paragraphs, tables, and footnotes from a DOCX file."""
    file_path = Path(data_dir) / filename
    document = Document(file_path)

    footnotes: List[str] = []
    try:
        with zipfile.ZipFile(file_path) as archive:
            xml_blob = archive.read("word/footnotes.xml")
        footnotes = _extract_footnotes_from_xml(xml_blob)
    except KeyError:
        footnotes = []

    tables: List[List[List[str]]] = []
    for table in document.tables:
        rows: List[List[str]] = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append(rows)

    return {
        "source_file": str(file_path),
        "paragraphs": [paragraph.text for paragraph in document.paragraphs],
        "tables": tables,
        "footnotes": footnotes,
    }


def normalize_extracted_content(raw: Dict[str, Any]) -> Dict[str, Any]:
    """Normalize whitespace and row labels from raw DOCX extraction."""
    normalized_tables: List[List[List[str]]] = []
    for table in raw.get("tables", []):
        normalized_rows: List[List[str]] = []
        for row in table:
            normalized_rows.append(
                [
                    _normalize_row_label(cell) if idx == 0 else _normalize_whitespace(cell)
                    for idx, cell in enumerate(row)
                ]
            )
        normalized_tables.append(normalized_rows)

    return {
        "source_file": raw.get("source_file", ""),
        "paragraphs": [
            value
            for value in (_normalize_whitespace(paragraph) for paragraph in raw.get("paragraphs", []))
            if value
        ],
        "tables": normalized_tables,
        "footnotes": [
            value
            for value in (_normalize_whitespace(note) for note in raw.get("footnotes", []))
            if value
        ],
    }


def detect_summary_statistics(rows: List[Dict[str, Any]]) -> List[str]:
    """Infer requested summary statistics from row labels."""
    labels = " | ".join(_normalize_whitespace(row.get("row_label", "")).lower() for row in rows)
    statistics: List[str] = []
    if re.search(r"(^|[| ])n($|[| ])", labels):
        statistics.append("count")
    if "mean (sd)" in labels:
        statistics.extend(["mean", "sd"])
    if "median" in labels:
        statistics.append("median")
    if "min-max" in labels or "min–max" in labels or "minmax" in labels:
        statistics.extend(["min", "max"])
    return statistics

